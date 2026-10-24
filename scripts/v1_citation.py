#!/usr/bin/env python3
"""v1_citation.py — V1-b 引用回填（verification-plan §3.1 第二段）

10 条"条款→证据→正文引用"链的导出与回查：
  链 = 条款定位器（A 源 DocxParagraph/DocxTableCell）+ 证据定位器（EV/B1 源
       PdfPageText）+ 正文引用定位器（导出件 DocxParagraph + CIT-书签锚）

三级断言（引擎语义仿真，ooxml_util——SEMUL 口径见该文件头注）：
  R1 定位回查：导出后重开文件，重算段落指纹表，10 条链全部 valid
     （位置在 + 指纹同——两级判定）
  R2 内容失效：改写一条引用段（run 级改写，书签保留）→ 恰该链 stale_content，
     其余 valid（"位置在但内容换"二级失效 + 只降不升：降级不自发恢复）
  R3 结构漂移：删除更早的正文段 → 索引型定位器全体漂移失效（stale_content /
     stale_position）；CIT- 书签锚仍可再寻址（引擎 ByAnchor 语义的仿真）→
     显式重验通道（reverify_claim 语义）重绑后恢复 valid

R3 即 doc-engine-audit §2.1 结论的运行时实证：裸索引定位器在结构变化下脆弱，
锚点（书签/BlockId）+ 显式重验是恢复通道——静默重绑被禁（引擎 degrade_linkage
只降不升纪律）。
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from ooxml_util import emit_docx_paragraph, fnv64, norm_text, resolve_docx_paragraph  # noqa: E402
import v1_export  # noqa: E402

REPO = Path(__file__).resolve().parent.parent


# ---------------------------------------------------------------- 书签解析
def bookmark_paragraph_map(docx_path: Path) -> dict[str, int]:
    """CIT- 书签名 → 正文段落序（与 python-docx doc.paragraphs 同口径：
    w:body 直接子 w:p 顺序计数）。"""
    doc = Document(str(docx_path))
    body = doc.element.body
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    result = {}
    idx = -1
    for child in body:
        if child.tag == qn("w:p"):
            idx += 1
            for bs in child.findall(".//w:bookmarkStart", ns):
                name = bs.get(qn("w:name"))
                if name and name.startswith("CIT-"):
                    result[name] = idx
    return result


def paragraph_fingerprint_table(docx_path: Path) -> list[str]:
    """正文段落规范化文本指纹表（引擎 refresh_document 的 paragraph_fps 入参）。"""
    doc = Document(str(docx_path))
    return [fnv64(norm_text(p.text)) for p in doc.paragraphs]


def paragraph_texts(docx_path: Path) -> list[str]:
    doc = Document(str(docx_path))
    return [p.text for p in doc.paragraphs]


# ---------------------------------------------------------------- 主流程
def run_citation_checks(bright_docx: Path, workdir: Path) -> dict:
    chains_src = v1_export.build_chain_locators()          # 条款/证据侧定位器
    bm = bookmark_paragraph_map(bright_docx)
    texts = paragraph_texts(bright_docx)
    chains = []
    for c in chains_src:
        gt = c["gt_id"]
        name = f"CIT-{gt}"
        pi = bm[name]
        chains.append({**c, "bookmark": name, "body_paragraph_index": pi,
                       "body_locator": emit_docx_paragraph(pi, texts[pi])})

    # ---- R1 定位回查（导出后重开）
    fps = paragraph_fingerprint_table(bright_docx)
    r1 = {c["gt_id"]: resolve_docx_paragraph(c["body_locator"], fps) for c in chains}
    r1_ok = all(v == "valid" for v in r1.values())

    # ---- R2 内容失效（run 级改写 GT-01 段，书签保留）
    mod2 = workdir / "modified_content.docx"
    shutil.copy(bright_docx, mod2)
    doc = Document(str(mod2))
    p = doc.paragraphs[bm["CIT-GT-01"]]
    p.runs[0].text += "（补充说明：证照在有效期内。）"
    doc.save(str(mod2))
    fps2 = paragraph_fingerprint_table(mod2)
    r2 = {c["gt_id"]: resolve_docx_paragraph(c["body_locator"], fps2) for c in chains}
    r2_stale_expected = {"GT-01"}
    r2_ok = (r2["GT-01"] == "stale_content"
             and {k for k, v in r2.items() if v != "valid"} == r2_stale_expected)

    # ---- R3 结构漂移（删除资格响应章引言段——位于全部引用段之前）
    mod3 = workdir / "modified_structure.docx"
    shutil.copy(bright_docx, mod3)
    doc = Document(str(mod3))
    intro_idx = next(i for i, t in enumerate(paragraph_texts(bright_docx))
                     if t.startswith("本章按招标文件第三章资格要求逐条响应"))
    doc.paragraphs[intro_idx]._p.getparent().remove(doc.paragraphs[intro_idx]._p)
    doc.save(str(mod3))
    fps3 = paragraph_fingerprint_table(mod3)
    r3 = {c["gt_id"]: resolve_docx_paragraph(c["body_locator"], fps3) for c in chains}
    r3_stale = {k for k, v in r3.items() if v != "valid"}
    r3_ok = r3_stale == {c["gt_id"] for c in chains}  # 全部在删除点之后 → 全体漂移失效

    # ---- R3b 书签再寻址 + 显式重验恢复（reverify_claim 语义）
    bm3 = bookmark_paragraph_map(mod3)
    rebind = {}
    for c in chains:
        gt = c["gt_id"]
        new_idx = bm3.get(f"CIT-{gt}")
        texts3 = paragraph_texts(mod3)
        if new_idx is None:
            rebind[gt] = "bookmark_lost"
            continue
        new_loc = emit_docx_paragraph(new_idx, texts3[new_idx])
        status_after = resolve_docx_paragraph(new_loc, fps3)
        rebind[gt] = status_after
    r3b_ok = all(v == "valid" for v in rebind.values())

    return {
        "chains": [{k: c[k] for k in ("gt_id", "clause_no", "bookmark",
                                      "body_paragraph_index", "clause_locator",
                                      "evidence_file", "evidence_locator")}
                   for c in chains],
        "R1_recheck": {"results": r1, "pass": r1_ok},
        "R2_content_stale": {"results": r2, "pass": r2_ok,
                             "note": "run 级改写（书签保留）：恰 GT-01 stale_content，其余 valid"},
        "R3_structure_drift": {"results": r3, "pass": r3_ok,
                               "deleted_paragraph_index": intro_idx,
                               "note": "删除更早段落后索引型定位器全体漂移失效（stale）"},
        "R3b_bookmark_rebind": {"results": rebind, "pass": r3b_ok,
                                "note": "CIT- 书签再寻址 + 显式重验（reverify_claim 语义）恢复 valid"},
    }


def main() -> int:
    out = Path(sys.argv[sys.argv.index("--out") + 1]) if "--out" in sys.argv else REPO / "artifacts" / "v1"
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        bright = out / "技术标-导出样本.docx"
        v1_export.build_bid_docx(bright, "bright", workdir=td)
        report = run_citation_checks(bright, td)
        (out / "citation-report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        ok = report["R1_recheck"]["pass"] and report["R2_content_stale"]["pass"] \
            and report["R3_structure_drift"]["pass"] and report["R3b_bookmark_rebind"]["pass"]
        print("[V1-b] R1 定位回查 10/10:", report["R1_recheck"]["pass"])
        print("[V1-b] R2 内容失效恰 1 条:", report["R2_content_stale"]["pass"])
        print("[V1-b] R3 结构漂移全体失效:", report["R3_structure_drift"]["pass"])
        print("[V1-b] R3b 书签重验恢复:", report["R3b_bookmark_rebind"]["pass"])
        return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
